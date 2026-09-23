import openpyxl
import json
import os
import re

excel_path = os.path.abspath("data/VOLAMP_Product_MASTER_DATA.xlsx")
output_dir = os.path.abspath("server/data")
os.makedirs(output_dir, exist_ok=True)
output_file = os.path.join(output_dir, "products_db.json")

print(f"Loading workbook from {excel_path}...")
wb = openpyxl.load_workbook(excel_path, data_only=True)

all_products = []

def clean_val(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s != "" and s != "-" and s != "None" else None

def parse_num_price(val_str):
    if not val_str:
        return None
    cleaned = re.sub(r"[^\d.]", "", str(val_str))
    try:
        f = float(cleaned)
        return int(round(f))
    except:
        return None

# ==========================================
# 1. WIRE CABLES (2,856 items)
# ==========================================
sheet1 = wb["WIRE  CABLES "]
rows1 = list(sheet1.iter_rows(values_only=True))
headers1 = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(rows1[0])]
h_map1 = {h: i for i, h in enumerate(headers1)}

for idx, r in enumerate(rows1[1:], start=1):
    if not any(c is not None and str(c).strip() != "" for c in r):
        continue
    
    brand = clean_val(r[h_map1.get("Brand", 7)]) or "Volamp"
    prod_cat = clean_val(r[h_map1.get("Product Category", 2)]) or "Wires & Cables"
    sub_cat = clean_val(r[h_map1.get("Sub-Category", 3)]) or "Cables"
    size = clean_val(r[h_map1.get("Size - Sq Mm", 18)])
    material = clean_val(r[h_map1.get("Conductor Material", 12)]) or "Copper"
    unit = clean_val(r[h_map1.get("Unit / Packing", 6)]) or "Per Meter"
    rate = clean_val(r[h_map1.get("Rate (Rs.)", 10)])
    disc = clean_val(r[h_map1.get("Discount Percent", 9)])
    disc_rate = clean_val(r[h_map1.get("Disscounted rate", 8)])
    avail = clean_val(r[h_map1.get("Availablity", 11)]) or "IN STOCK"
    image_3d = clean_val(r[h_map1.get("3D IMAGE OF THE WIRE", 25)])
    prod_name = clean_val(r[h_map1.get("Product Name", 24)])
    
    if not prod_name:
        prod_name = f"{size or ''} SQMM {material} Cable {sub_cat}".strip()
    
    product_id = f"CAB-{idx:06d}"
    sku = f"SKU-{product_id}-{brand.upper().replace(' ', '')}"
    
    specs = {
        "conductorConstruction": clean_val(r[h_map1.get("Conductor Construction (No./Dia mm)", 4)]),
        "currentRatingAmp": clean_val(r[h_map1.get("Current Rating (Amp)", 5)]),
        "typeOfArmour": clean_val(r[h_map1.get("Type Of Armour", 13)]),
        "color": clean_val(r[h_map1.get("Color", 14)]),
        "conductorClass": clean_val(r[h_map1.get("Conductor Class", 15)]),
        "innerSheathMaterial": clean_val(r[h_map1.get("Inner Sheath Material", 16)]),
        "insulationType": clean_val(r[h_map1.get("Insulation Type", 17)]),
        "shieldingType": clean_val(r[h_map1.get("Sheilding Type", 19)]),
        "voltageRating": clean_val(r[h_map1.get("Voltage Rating", 20)]),
        "cores": clean_val(r[h_map1.get("Core - Pair - Triad - Quad", 21)]),
        "warranty": clean_val(r[h_map1.get("Warranty", 22)]),
        "outerSheathMaterial": clean_val(r[h_map1.get("Outer Sheath Material", 23)]),
        "cableType": clean_val(r[h_map1.get("TYPE", 1)])
    }
    
    all_products.append({
        "productId": product_id,
        "sku": sku,
        "name": prod_name,
        "brand": brand,
        "category": "Wires & Cables",
        "subcategory": sub_cat,
        "description": f"{brand} {prod_name} - {prod_cat}. Rated {specs['voltageRating'] or '1100V'}, {specs['cores'] or ''} {specs['typeOfArmour'] or ''} conductor.",
        "size": f"{size} SQMM" if size else None,
        "material": material,
        "unit": unit,
        "price": f"₹{float(rate):.2f}" if parse_num_price(rate) is not None else (rate or "On Request"),
        "numericPrice": parse_num_price(disc_rate or rate),
        "discount": disc or "40 %",
        "discountedPrice": f"₹{float(disc_rate):.2f}" if parse_num_price(disc_rate) is not None else None,
        "currency": "INR",
        "availability": avail,
        "moq": "100 Meters / Coil",
        "imageUrl": image_3d,
        "threeDImageUrl": image_3d,
        "productUrl": None,
        "status": "Active",
        "specifications": json.dumps(specs),
        "sheetSource": "WIRE CABLES",
        "lastUpdated": "2026-09-22T00:00:00.000Z"
    })

print(f"Processed Wires & Cables: {len(all_products)} products")

# ==========================================
# 2. SWITCH GEARS (221 items)
# ==========================================
sheet2 = wb["switch gears "]
rows2 = list(sheet2.iter_rows(values_only=True))
headers2 = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(rows2[0])]
h_map2 = {h: i for i, h in enumerate(headers2)}

swg_start_len = len(all_products)
for idx, r in enumerate(rows2[1:], start=1):
    if not any(c is not None and str(c).strip() != "" for c in r):
        continue
    
    brand = clean_val(r[h_map2.get("BRAND", 11)]) or "Volamp"
    prod_cat = clean_val(r[h_map2.get("Product_Category", 1)]) or "Switching & Control"
    sub_cat = clean_val(r[h_map2.get("Subcategory", 2)]) or "Switchgear"
    desc = clean_val(r[h_map2.get("Product_Description", 3)])
    poles = clean_val(r[h_map2.get("Poles_Phase", 5)])
    current = clean_val(r[h_map2.get("Current_Rating_A", 6)])
    voltage = clean_val(r[h_map2.get("Voltage_Rating_V", 7)])
    img = clean_val(r[h_map2.get("images", 12)])
    rate = clean_val(r[h_map2.get("Rate (Rs.)", 19)])
    disc = clean_val(r[h_map2.get("Discount", 18)])
    disc_rate = clean_val(r[h_map2.get("Discounted Rate", 17)])
    unit = clean_val(r[h_map2.get("Pricing", 20)]) or "Per Unit"
    prod_name = clean_val(r[h_map2.get("Product Name", 22)]) or desc or f"{sub_cat} {poles or ''} {current or ''}A"
    
    product_id = f"SWG-{idx:06d}"
    sku = f"SKU-{product_id}-{brand.upper().replace(' ', '')}"
    
    specs = {
        "hsnCode": clean_val(r[h_map2.get("HSN_Code", 4)]),
        "polesPhase": poles,
        "currentRatingA": current,
        "voltageRatingV": voltage,
        "breakingCapacityKa": clean_val(r[h_map2.get("Breaking_Capacity_kA", 8)]),
        "tripCurve": clean_val(r[h_map2.get("Trip Characteristic_Curve", 9)]),
        "sensitivityMa": clean_val(r[h_map2.get("Sensitivity_mA", 10)]),
        "executionMounting": clean_val(r[h_map2.get("Execution_Mounting", 13)]),
        "auxiliaryContacts": clean_val(r[h_map2.get("Auxiliary_Contacts", 14)]),
        "frameSize": clean_val(r[h_map2.get("Frame_Size", 15)]),
        "notesSpecifications": clean_val(r[h_map2.get("Notes_Specifications", 21)])
    }
    
    all_products.append({
        "productId": product_id,
        "sku": sku,
        "name": prod_name,
        "brand": brand,
        "category": "Switchgear",
        "subcategory": sub_cat,
        "description": desc or f"{brand} {prod_name}. {specs['notesSpecifications'] or ''}",
        "size": f"{current}A" if current else None,
        "material": "Polycarbonate / Silver Alloy Contacts",
        "unit": unit,
        "price": f"₹{float(rate):.2f}" if parse_num_price(rate) is not None else (rate or "On Request"),
        "numericPrice": parse_num_price(disc_rate or rate),
        "discount": disc or "40 %",
        "discountedPrice": f"₹{float(disc_rate):.2f}" if parse_num_price(disc_rate) is not None else None,
        "currency": "INR",
        "availability": "IN STOCK",
        "moq": "1 Unit",
        "imageUrl": img,
        "threeDImageUrl": img,
        "productUrl": None,
        "status": "Active",
        "specifications": json.dumps(specs),
        "sheetSource": "switch gears",
        "lastUpdated": "2026-09-22T00:00:00.000Z"
    })

print(f"Processed Switchgear: {len(all_products) - swg_start_len} products")

# ==========================================
# 3. LUGS (14 items)
# ==========================================
sheet3 = wb["Lugs "]
rows3 = list(sheet3.iter_rows(values_only=True))
headers3 = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(rows3[0])]
h_map3 = {h: i for i, h in enumerate(headers3)}

lug_start_len = len(all_products)
for idx, r in enumerate(rows3[1:], start=1):
    if not any(c is not None and str(c).strip() != "" for c in r):
        continue
    
    mat = clean_val(r[h_map3.get("Material", 1)]) or "Copper"
    cat = clean_val(r[h_map3.get("Category", 2)]) or "Ring"
    sub_cat = clean_val(r[h_map3.get("sub  category", 3)]) or "Standard"
    prod_name = clean_val(r[h_map3.get("Product Name", 4)]) or f"{cat} Type {mat} Lugs {sub_cat}"
    rate = clean_val(r[h_map3.get("Rate (Rs.)", 5)]) or "On request"
    disc = clean_val(r[h_map3.get("Discount", 6)]) or "On request"
    size = clean_val(r[h_map3.get("size", 7)]) or "All Sizes Available"
    img = clean_val(r[h_map3.get("images", 8)])
    
    product_id = f"LUG-{idx:06d}"
    sku = f"SKU-{product_id}-VOLAMP"
    
    specs = {
        "material": mat,
        "lugType": cat,
        "barrelStyle": sub_cat,
        "sizeRange": size,
        "standard": "IS 8309 / DIN 46235"
    }
    
    all_products.append({
        "productId": product_id,
        "sku": sku,
        "name": prod_name,
        "brand": "Volamp",
        "category": "Lugs",
        "subcategory": cat,
        "description": f"Heavy duty electrolytic {mat} crimping cable lugs. Designed for high conductivity and vibration resistance.",
        "size": size,
        "material": mat,
        "unit": "Per 100 Pcs",
        "price": rate,
        "numericPrice": parse_num_price(rate) or 150,
        "discount": disc,
        "discountedPrice": None,
        "currency": "INR",
        "availability": "IN STOCK",
        "moq": "100 Pcs",
        "imageUrl": img,
        "threeDImageUrl": img,
        "productUrl": None,
        "status": "Active",
        "specifications": json.dumps(specs),
        "sheetSource": "Lugs",
        "lastUpdated": "2026-09-22T00:00:00.000Z"
    })

print(f"Processed Lugs: {len(all_products) - lug_start_len} products")

# ==========================================
# 4. PVC PIPE (93 items)
# ==========================================
sheet4 = wb["Pvc Pipe"]
rows4 = list(sheet4.iter_rows(values_only=True))
headers4 = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(rows4[0])]
h_map4 = {h: i for i, h in enumerate(headers4)}

pvc_start_len = len(all_products)
for idx, r in enumerate(rows4[1:], start=1):
    if not any(c is not None and str(c).strip() != "" for c in r):
        continue
    
    family = clean_val(r[h_map4.get("Product_Family", 0)]) or "PVC Conduits"
    size = clean_val(r[h_map4.get("Size", 1)])
    std_pack = clean_val(r[h_map4.get("Standard_Pack", 2)])
    rate = clean_val(r[h_map4.get("Rate (Rs.)", 3)])
    unit = clean_val(r[h_map4.get("Unit", 4)]) or "₹/m"
    disc = clean_val(r[h_map4.get("discount", 5)]) or "20 %"
    disc_rate = clean_val(r[h_map4.get("Discounted Rate", 6)])
    color = clean_val(r[h_map4.get("COLOR", 7)]) or "Ivory / White"
    spec_text = clean_val(r[h_map4.get("Specifications", 8)])
    prod_name = clean_val(r[h_map4.get("Porduct name", 9)]) or f"{std_pack or ''} {family} {size or ''}mm PVC PIPE".strip()
    img = clean_val(r[h_map4.get("Images", 10)])
    
    product_id = f"PVC-{idx:06d}"
    sku = f"SKU-{product_id}-VOLAMP"
    
    specs = {
        "mechanicalStress": family,
        "color": color,
        "standardPack": std_pack,
        "specifications": spec_text,
        "standard": "IS 9537 Part 3"
    }
    
    all_products.append({
        "productId": product_id,
        "sku": sku,
        "name": prod_name,
        "brand": "Volamp",
        "category": "PVC Pipe",
        "subcategory": family.split("(")[0].strip() if "(" in family else family,
        "description": spec_text or f"High impact rigid {family} electrical PVC conduit pipe in {color} color. Smooth internal bore for easy wire pulling.",
        "size": f"{size} mm" if size else None,
        "material": "Rigid Unplasticized PVC (uPVC)",
        "unit": unit,
        "price": f"₹{float(rate):.2f}" if parse_num_price(rate) is not None else (rate or "On Request"),
        "numericPrice": parse_num_price(disc_rate or rate),
        "discount": disc,
        "discountedPrice": f"₹{float(disc_rate):.2f}" if parse_num_price(disc_rate) is not None else None,
        "currency": "INR",
        "availability": "IN STOCK",
        "moq": std_pack or "50 Pcs x 3m",
        "imageUrl": img,
        "threeDImageUrl": img,
        "productUrl": None,
        "status": "Active",
        "specifications": json.dumps(specs),
        "sheetSource": "Pvc Pipe",
        "lastUpdated": "2026-09-22T00:00:00.000Z"
    })

print(f"Processed PVC Pipe: {len(all_products) - pvc_start_len} products")

# ==========================================
# 5. GLANDS (90 items)
# ==========================================
sheet5 = wb["Glands "]
rows5 = list(sheet5.iter_rows(values_only=True))
headers5 = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(rows5[0])]
h_map5 = {h: i for i, h in enumerate(headers5)}

gld_start_len = len(all_products)
for idx, r in enumerate(rows5[1:], start=1):
    if not any(c is not None and str(c).strip() != "" for c in r):
        continue
    
    cat = clean_val(r[h_map5.get("Category", 1)]) or "Single Compression"
    size_code = clean_val(r[h_map5.get("Size Code", 2)])
    size_inch = clean_val(r[h_map5.get("Size (Inch)", 3)])
    cable_od = clean_val(r[h_map5.get("Suitable Cable OD (mm)", 4)])
    nipple = clean_val(r[h_map5.get("Nipple Thread", 5)])
    rate = clean_val(r[h_map5.get("Rate (Rs.)", 6)])
    box_pcs = clean_val(r[h_map5.get("Packing per Box (pcs)", 7)])
    box_price = clean_val(r[h_map5.get("Price per Box ()", 8)])
    shroud_price = clean_val(r[h_map5.get("PVC Shroud Price ()", 9)])
    earth_tag_price = clean_val(r[h_map5.get("Brass Earth Tag Price ()", 10)])
    prod_name = clean_val(r[h_map5.get("Product name", 11)]) or f"{cat} {size_code or ''} CABLE GLAND".strip()
    full_set_price = clean_val(r[h_map5.get("Full Set Price\n (Gland+Shroud+Tag) ()", 12)])
    img_3d = clean_val(r[h_map5.get("3D IMAGE", 13)])
    
    product_id = f"GLD-{idx:06d}"
    sku = f"SKU-{product_id}-VOLAMP"
    
    specs = {
        "compressionType": cat,
        "sizeCode": size_code,
        "sizeInch": size_inch,
        "suitableCableOdMm": cable_od,
        "nippleThread": nipple,
        "packingPerBoxPcs": box_pcs,
        "boxPrice": box_price,
        "shroudPrice": shroud_price,
        "earthTagPrice": earth_tag_price,
        "fullSetPrice": full_set_price,
        "standard": "BS 6121 / IS 12943"
    }
    
    numeric_p = parse_num_price(rate)
    
    all_products.append({
        "productId": product_id,
        "sku": sku,
        "name": prod_name,
        "brand": "Volamp",
        "category": "Glands",
        "subcategory": cat,
        "description": f"Nickel plated brass {cat} cable gland. Suitable for cable OD {cable_od or 'all sizes'} mm with {nipple or 'standard'} nipple thread.",
        "size": size_code or size_inch,
        "material": "Brass (Nickel Plated)",
        "unit": "Per Piece",
        "price": f"₹{numeric_p:.2f}" if numeric_p is not None else (rate or "On Request"),
        "numericPrice": numeric_p,
        "discount": "Standard Wholesale",
        "discountedPrice": None,
        "currency": "INR",
        "availability": "IN STOCK",
        "moq": f"{box_pcs} Pcs / Box" if box_pcs else "10 Pcs",
        "imageUrl": img_3d,
        "threeDImageUrl": img_3d,
        "productUrl": None,
        "status": "Active",
        "specifications": json.dumps(specs),
        "sheetSource": "Glands",
        "lastUpdated": "2026-09-22T00:00:00.000Z"
    })

print(f"Processed Glands: {len(all_products) - gld_start_len} products")

# ==========================================
# 6. WIRING DEVICE (14 items)
# ==========================================
sheet6 = wb["Wiring device"]
rows6 = list(sheet6.iter_rows(values_only=True))
headers6 = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(rows6[0])]
h_map6 = {h: i for i, h in enumerate(headers6)}

wd_start_len = len(all_products)
for idx, r in enumerate(rows6[1:], start=1):
    if not any(c is not None and str(c).strip() != "" for c in r):
        continue
    
    item_id = clean_val(r[h_map6.get("Item_ID", 0)]) or f"WD{idx:03d}"
    cat = clean_val(r[h_map6.get("Category", 1)]) or "Wiring Accessory"
    sub_cat = clean_val(r[h_map6.get("Sub - Category", 2)]) or "Wiring Devices"
    spec = clean_val(r[h_map6.get("Specification", 3)])
    mat = clean_val(r[h_map6.get("Material", 4)]) or "Polycarbonate / PVC"
    volt = clean_val(r[h_map6.get("Voltage_Rating", 5)])
    curr = clean_val(r[h_map6.get("Current_Rating", 6)])
    unit = clean_val(r[h_map6.get("Unit", 10)]) or "Piece"
    rate = clean_val(r[h_map6.get("Rate (Rs.)", 11)]) or "ON REQUEST"
    disc = clean_val(r[h_map6.get("DISCOUNT", 12)]) or "0 %"
    app = clean_val(r[h_map6.get("Application", 13)])
    img = clean_val(r[h_map6.get("images", 14)])
    
    # Standardize or preserve ID
    product_id = f"WD-{idx:06d}" if not item_id.startswith("WD") else f"WD-{int(re.sub(r'[^0-9]', '', item_id) or idx):06d}"
    sku = f"SKU-{product_id}-VOLAMP"
    prod_name = f"{cat} - {spec}" if spec else f"{cat} ({sub_cat})"
    
    specs = {
        "voltageRating": volt,
        "currentRating": curr,
        "widthOrDiameter": clean_val(r[h_map6.get("Width_or_Diameter", 7)]),
        "length": clean_val(r[h_map6.get("Length", 8)]),
        "height": clean_val(r[h_map6.get("Height", 9)]),
        "application": app,
        "originalItemId": item_id
    }
    
    all_products.append({
        "productId": product_id,
        "sku": sku,
        "name": prod_name,
        "brand": "Volamp",
        "category": "Wiring Devices",
        "subcategory": sub_cat,
        "description": f"Industrial-grade electrical {cat}. Designed for {app or 'general electrical installations'}. Material: {mat}.",
        "size": f"{curr}" if curr and curr != "N/A" else None,
        "material": mat,
        "unit": unit,
        "price": rate,
        "numericPrice": parse_num_price(rate) or 250,
        "discount": disc,
        "discountedPrice": None,
        "currency": "INR",
        "availability": "IN STOCK",
        "moq": "10 Pcs",
        "imageUrl": img,
        "threeDImageUrl": img,
        "productUrl": None,
        "status": "Active",
        "specifications": json.dumps(specs),
        "sheetSource": "Wiring device",
        "lastUpdated": "2026-09-22T00:00:00.000Z"
    })

print(f"Processed Wiring Devices: {len(all_products) - wd_start_len} products")

# ==========================================
# 7. EARTHING WIRES (37 items)
# ==========================================
sheet7 = wb["Earthing Wires"]
rows7 = list(sheet7.iter_rows(values_only=True))
headers7 = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(rows7[0])]
h_map7 = {h: i for i, h in enumerate(headers7)}

ear_start_len = len(all_products)
for idx, r in enumerate(rows7[1:], start=1):
    if not any(c is not None and str(c).strip() != "" for c in r):
        continue
    
    raw_id = clean_val(r[h_map7.get("Product ID", 0)]) or f"EAR-{idx:03d}"
    prod_cat = clean_val(r[h_map7.get("Product Category", 1)]) or "Earthing Materials"
    sub_cat = clean_val(r[h_map7.get("Subcategory", 2)]) or "Electrodes"
    prod_name = clean_val(r[h_map7.get("Product Name", 3)]) or f"Earthing Product {raw_id}"
    mat = clean_val(r[h_map7.get("Material", 4)]) or "Copper Bonded"
    key_spec = clean_val(r[h_map7.get("Key Specifications", 5)])
    size_range = clean_val(r[h_map7.get("Size / Range", 6)])
    unit = clean_val(r[h_map7.get("Selling Unit", 7)]) or "Piece"
    app = clean_val(r[h_map7.get("Primary Application", 8)])
    add_spec = clean_val(r[h_map7.get("Specifications ( Additional )", 9)])
    moq = clean_val(r[h_map7.get("MOQ", 10)])
    disc_rates = clean_val(r[h_map7.get("Discounted rates", 11)])
    disc = clean_val(r[h_map7.get("Discount", 12)])
    rate = clean_val(r[h_map7.get("Rate (Rs.)", 13)])
    img = clean_val(r[h_map7.get("Image URL", 14)])
    seo = clean_val(r[h_map7.get("SEO Keywords", 15)])
    
    # Preserve original ID e.g. EAR-001
    product_id = raw_id
    sku = f"SKU-{product_id}-VOLAMP"
    
    specs = {
        "keySpecifications": key_spec,
        "additionalSpecifications": add_spec,
        "primaryApplication": app,
        "seoKeywords": seo,
        "standard": "IS 3043 / IEEE 80"
    }
    
    num_p = parse_num_price(disc_rates or rate)
    
    all_products.append({
        "productId": product_id,
        "sku": sku,
        "name": prod_name,
        "brand": "Volamp",
        "category": "Earthing Wires",
        "subcategory": sub_cat,
        "description": key_spec or f"Maintenance-free {prod_name}. Suitable for {app or 'electrical grounding'}. Material: {mat}.",
        "size": size_range,
        "material": mat,
        "unit": unit,
        "price": f"₹{float(rate):.2f}" if parse_num_price(rate) is not None else (rate or "On Request"),
        "numericPrice": num_p,
        "discount": disc or "0 %",
        "discountedPrice": f"₹{float(disc_rates):.2f}" if parse_num_price(disc_rates) is not None else None,
        "currency": "INR",
        "availability": "IN STOCK",
        "moq": f"{int(float(moq))} {unit}" if moq and parse_num_price(moq) else (moq or "5 Pieces"),
        "imageUrl": img,
        "threeDImageUrl": img,
        "productUrl": None,
        "status": "Active",
        "specifications": json.dumps(specs),
        "sheetSource": "Earthing Wires",
        "lastUpdated": "2026-09-22T00:00:00.000Z"
    })

print(f"Processed Earthing Wires: {len(all_products) - ear_start_len} products")

# ==========================================
# 8. SOLAR (60 items)
# ==========================================
sheet8 = wb["Solar"]
rows8 = list(sheet8.iter_rows(values_only=True))
headers8 = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(rows8[0])]
h_map8 = {h: i for i, h in enumerate(headers8)}

sol_start_len = len(all_products)
for idx, r in enumerate(rows8[1:], start=1):
    if not any(c is not None and str(c).strip() != "" for c in r):
        continue
    
    cat = clean_val(r[h_map8.get("Category", 0)]) or "Solar Cable"
    sub_cat = clean_val(r[h_map8.get("Sub category", 1)]) or "PV1-F"
    size = clean_val(r[h_map8.get("Size_Sqmm", 2)])
    watt = clean_val(r[h_map8.get("Watt", 3)])
    pack_mtrs = clean_val(r[h_map8.get("Standard_Packing_Mtrs", 4)])
    pack_pcs = clean_val(r[h_map8.get("Standard_Packing_Pcs", 6)])
    disc_rates = clean_val(r[h_map8.get("Discounted Rates", 7)]) or clean_val(r[h_map8.get("Discounted Rate", 5)])
    disc = clean_val(r[h_map8.get("Discount", 8)])
    rate = clean_val(r[h_map8.get("Rate (Rs.)", 9)])
    curr_60 = clean_val(r[h_map8.get("Current_60C_A", 10)])
    curr_40 = clean_val(r[h_map8.get("Current_40C_A", 13)])
    eff = clean_val(r[h_map8.get("Effeicency", 11)])
    war = clean_val(r[h_map8.get("Warrenty", 12)])
    prod_name = clean_val(r[h_map8.get("Product name", 14)]) or f"{sub_cat} {size or watt or ''} Solar Product".strip()
    img = clean_val(r[h_map8.get("images", 15)])
    
    product_id = f"SOL-{idx:06d}"
    sku = f"SKU-{product_id}-VOLAMP"
    
    specs = {
        "watt": watt,
        "standardPackingMtrs": pack_mtrs,
        "standardPackingPcs": pack_pcs,
        "current60CAmp": curr_60,
        "current40CAmp": curr_40,
        "efficiency": eff,
        "warranty": war,
        "standard": "EN 50618 / TÜV Certified"
    }
    
    num_p = parse_num_price(disc_rates or rate)
    
    all_products.append({
        "productId": product_id,
        "sku": sku,
        "name": prod_name,
        "brand": "Volamp",
        "category": "Solar",
        "subcategory": cat.strip() if cat else sub_cat,
        "description": f"TÜV certified photovoltaic {cat} ({sub_cat}). Rated for UV, ozone and weather resistance in solar power generation plants.",
        "size": f"{size} SQMM" if size else (f"{watt}W" if watt else None),
        "material": "Tinned Copper / Crosslinked Polyolefin",
        "unit": "Per Meter" if "cable" in cat.lower() else "Per Piece",
        "price": f"₹{float(rate):.2f}" if parse_num_price(rate) is not None else (rate or "On Request"),
        "numericPrice": num_p,
        "discount": f"{float(disc)*100:.0f} %" if disc and parse_num_price(disc) is not None and float(disc) < 1 else (disc or "30 %"),
        "discountedPrice": f"₹{float(disc_rates):.2f}" if parse_num_price(disc_rates) is not None else None,
        "currency": "INR",
        "availability": "IN STOCK",
        "moq": f"{int(float(pack_mtrs))} Mtrs" if pack_mtrs and parse_num_price(pack_mtrs) else "100 Mtrs",
        "imageUrl": img,
        "threeDImageUrl": img,
        "productUrl": None,
        "status": "Active",
        "specifications": json.dumps(specs),
        "sheetSource": "Solar",
        "lastUpdated": "2026-09-22T00:00:00.000Z"
    })

print(f"Processed Solar: {len(all_products) - sol_start_len} products")

print(f"\n==========================================")
print(f"TOTAL PRODUCTS INGESTED: {len(all_products)}")
print(f"Saving to {output_file}...")

with open(output_file, "w", encoding="utf-8") as f:
    json.dump(all_products, f, indent=2, ensure_ascii=False)

print(f"SUCCESS: Successfully saved {len(all_products)} products to {output_file}!")
